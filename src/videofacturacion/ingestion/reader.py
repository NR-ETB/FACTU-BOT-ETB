import csv
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from videofacturacion.config import Settings
from videofacturacion.domain.billing import reconcile
from videofacturacion.domain.models import Invoice, Issue, Line, SourceRef
from videofacturacion.domain.money import parse_money

HEADER = (
    "cliente_id",
    "nombre_cliente",
    "factura_id",
    "cuenta_id",
    "periodo_inicio",
    "periodo_fin",
    "fecha_vencimiento",
    "moneda",
    "total_fuente",
)
COLUMNS = HEADER + ("linea_id", "servicio", "concepto_codigo", "concepto_descripcion", "importe")
ETB_COLUMNS = (
    "id_cliente",
    "nombre_cliente",
    "numero_cuenta",
    "numero_factura",
    "tipo_servicio",
    "periodo_inicio",
    "periodo_fin",
    "fecha_emision",
    "fecha_vencimiento",
    "moneda",
    "valor_internet",
    "valor_television",
    "valor_telefonia",
    "subtotal_servicios",
    "descuento",
    "saldo_anterior",
    "otros_cargos",
    "total_pagar",
    "estado_factura",
)
ETB_MAP = {
    "cliente_id": "id_cliente",
    "cuenta_id": "numero_cuenta",
    "factura_id": "numero_factura",
    "total_fuente": "total_pagar",
}
PROFILES = {"conceptos_v1": COLUMNS, "etb_prueba_v1": ETB_COLUMNS}


class InputError(Exception):
    def __init__(self, issues: list[Issue]):
        self.issues = issues
        super().__init__(issues[0].message)


@dataclass
class ParsedImport:
    file_hash: str
    invoices: list[Invoice]


def fail(code: str, message: str, **location) -> None:
    raise InputError([Issue(code, message, **location)])


def cell_text(cell) -> str:
    value = cell.value
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return "" if value is None else str(value)


def read_table(data: bytes, filename: str, profile: str, settings: Settings):
    if profile not in PROFILES:
        fail("UNKNOWN_PROFILE", "Perfil de importación desconocido.")
    if not data or len(data) > settings.max_bytes:
        fail("FILE_SIZE", "El archivo está vacío o excede el tamaño permitido.")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            text = data.decode("utf-8-sig", errors="strict")
            reader = csv.reader(
                io.StringIO(text, newline=""),
                delimiter=";" if profile == "conceptos_v1" else ",",
                strict=True,
            )
            headings = next(reader)
            rows = []
            for row in reader:
                if not row or all(not x.strip() for x in row):
                    continue
                if len(rows) >= settings.max_lines:
                    fail("LINE_LIMIT", "El archivo excede el máximo de filas.")
                rows.append((reader.line_num, row))
        except (UnicodeError, csv.Error, StopIteration):
            fail("INVALID_CSV", "CSV inválido: revise UTF-8, comillas y delimitador del perfil.")
        sheet = "CSV"
    elif suffix == ".xlsx":
        # Limitar la expansión del contenedor antes de procesar XML.
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                    fail("XLSX_EXPANSION", "El contenido descomprimido supera el límite.")
                if len(archive.infolist()) > 1000:
                    fail("XLSX_EXPANSION", "El libro contiene demasiados componentes.")
            workbook = load_workbook(
                io.BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            try:
                if "facturacion" not in workbook.sheetnames:
                    fail("MISSING_SHEET", "El libro debe incluir la hoja facturacion.")
                worksheet = workbook["facturacion"]
                if worksheet.max_row > settings.max_lines + 1:
                    fail("LINE_LIMIT", "El libro excede el máximo de filas físicas.")
                if worksheet.max_column > len(PROFILES[profile]):
                    fail("INVALID_COLUMNS", "Hay columnas adicionales no admitidas.")
                iterator = worksheet.iter_rows()
                headings = [cell_text(c) for c in next(iterator)]
                rows = []
                for row_number, cells in enumerate(iterator, start=2):
                    if all(c.value is None for c in cells):
                        continue
                    if len(rows) >= settings.max_lines or row_number > settings.max_lines + 1:
                        fail("LINE_LIMIT", "El archivo excede el máximo de filas.")
                    for column, cell in enumerate(cells):
                        if cell.data_type == "f":
                            fail(
                                "FORMULA_NOT_ALLOWED",
                                "No se admiten fórmulas.",
                                row=row_number,
                                field=headings[column],
                                sheet="facturacion",
                            )
                    rows.append((row_number, [cell_text(c) for c in cells]))
            finally:
                workbook.close()
        except InputError:
            raise
        except Exception:
            fail("INVALID_XLSX", "No fue posible leer el libro XLSX.")
        sheet = "facturacion"
    else:
        fail("FILE_TYPE", "Solo se admiten archivos .csv y .xlsx.")
    if len(headings) != len(set(headings)) or set(headings) != set(PROFILES[profile]):
        fail("INVALID_COLUMNS", "Las columnas no corresponden exactamente al perfil elegido.")
    if not rows:
        fail("NO_ROWS", "El archivo no contiene facturas.")
    if any(len(values) != len(headings) for _, values in rows):
        fail("ROW_WIDTH", "Hay filas con un número de columnas diferente a la cabecera.")
    return sheet, [(n, dict(zip(headings, values, strict=True))) for n, values in rows]


def parse_file(data: bytes, filename: str, profile: str, settings: Settings) -> ParsedImport:
    sheet, rows = read_table(data, filename, profile, settings)
    digest = sha256(data).hexdigest()
    grouped: dict[str, Invoice] = {}
    for row_number, raw in rows:
        original = raw
        raw = {k: v.strip() for k, v in raw.items()}
        mapped = (
            {key: raw[ETB_MAP.get(key, key)] for key in HEADER}
            if profile == "etb_prueba_v1"
            else raw
        )
        key = f"{mapped['cuenta_id']}:{mapped['factura_id']}"
        refs = {
            k: SourceRef(
                digest,
                sheet,
                row_number,
                ETB_MAP.get(k, k) if profile == "etb_prueba_v1" else k,
                original[ETB_MAP.get(k, k)] if profile == "etb_prueba_v1" else original[k],
            )
            for k in HEADER
        }
        header = {key: mapped[key] for key in HEADER}
        invoice = grouped.setdefault(key, Invoice(key, header, refs))

        def issue(code, message, field):
            invoice.issues.append(Issue(code, message, row_number, field, sheet, key))

        for name, value in raw.items():
            if not value:
                issue("REQUIRED_FIELD", "El campo es obligatorio.", name)
            elif len(value) > 300 or any(ord(c) < 32 for c in value):
                issue("INVALID_TEXT", "Texto demasiado largo o con caracteres de control.", name)
        for name in ("cliente_id", "cuenta_id", "factura_id"):
            if not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", header[name]):
                issue(
                    "INVALID_ID",
                    "Identificador inválido: use letras, números, guion o guion bajo.",
                    refs[name].field,
                )
        dates = {}
        for name in ("periodo_inicio", "periodo_fin", "fecha_vencimiento"):
            try:
                if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", header[name]):
                    raise ValueError
                dates[name] = date.fromisoformat(header[name])
            except ValueError:
                issue("INVALID_DATE", "Fecha inválida: use AAAA-MM-DD.", name)
        if len(dates) == 3 and (
            dates["periodo_inicio"] > dates["periodo_fin"]
            or dates["fecha_vencimiento"] < dates["periodo_fin"]
        ):
            issue(
                "DATE_ORDER",
                "El período o el vencimiento tienen un orden inválido.",
                "fecha_vencimiento",
            )
        if header["moneda"] != "COP":
            issue("INVALID_CURRENCY", "La moneda admitida es COP.", "moneda")
        try:
            header["total_fuente"] = format(parse_money(header["total_fuente"]), ".2f")
        except ValueError as error:
            issue("INVALID_MONEY", str(error), refs["total_fuente"].field)
        if profile == "etb_prueba_v1" and len(invoice.lines):
            issue(
                "DUPLICATE_INVOICE",
                "Este perfil solo admite una fila por factura.",
                "numero_factura",
            )
        if invoice.header != header:
            issue(
                "HEADER_CONFLICT",
                "La cabecera difiere entre filas de la misma factura.",
                "factura_id",
            )

        if profile == "conceptos_v1":
            try:
                amount = parse_money(raw["importe"])
                invoice.lines.append(
                    Line(
                        raw["linea_id"],
                        raw["servicio"],
                        raw["concepto_codigo"],
                        raw["concepto_descripcion"],
                        amount,
                        SourceRef(digest, sheet, row_number, "importe", original["importe"]),
                    )
                )
            except ValueError as error:
                issue("INVALID_MONEY", str(error), "importe")
        else:
            fields = {
                "valor_internet": ("INTERNET", "Internet"),
                "valor_television": ("TELEVISION", "Televisión"),
                "valor_telefonia": ("TELEFONIA", "Telefonía"),
                "descuento": ("DESCUENTO", "Descuento"),
                "saldo_anterior": ("SALDO_ANTERIOR", "Saldo anterior"),
                "otros_cargos": ("OTROS_CARGOS", "Otros cargos"),
            }
            parsed = {}
            for name in (*fields, "subtotal_servicios"):
                try:
                    parsed[name] = parse_money(raw[name])
                    if parsed[name] < 0:
                        issue("INVALID_SIGN", "Este perfil requiere importes no negativos.", name)
                except ValueError as error:
                    issue("INVALID_MONEY", str(error), name)
            if all(
                k in parsed
                for k in (
                    "valor_internet",
                    "valor_television",
                    "valor_telefonia",
                    "subtotal_servicios",
                )
            ):
                subtotal = sum(
                    (parsed[k] for k in ("valor_internet", "valor_television", "valor_telefonia")),
                    Decimal(0),
                )
                if subtotal != parsed["subtotal_servicios"]:
                    issue(
                        "SUBTOTAL_MISMATCH",
                        "El subtotal no coincide con los tres servicios.",
                        "subtotal_servicios",
                    )
            for name, (code, description) in fields.items():
                if name in parsed:
                    amount = -parsed[name] if code == "DESCUENTO" else parsed[name]
                    invoice.lines.append(
                        Line(
                            name,
                            description,
                            code,
                            description,
                            amount,
                            SourceRef(digest, sheet, row_number, name, original[name]),
                        )
                    )
            invoice.refs["subtotal_servicios"] = SourceRef(
                digest, sheet, row_number, "subtotal_servicios", original["subtotal_servicios"]
            )
    if len(grouped) > settings.max_invoices:
        fail("INVOICE_LIMIT", "El archivo excede el máximo de facturas.")
    for invoice in grouped.values():
        reconcile(invoice)
    return ParsedImport(digest, list(grouped.values()))
